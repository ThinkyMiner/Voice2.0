"""The one-slide summary (PLAN.md §3 Phase 5): slide/one_slide.pptx, plus its chart PNG in reports/.

Reads the store only (a read-only connection). The numbers come from the decision log of one
promoted run and one harm-stopped run of the same design (scenario "win" / "harm" preferred),
their sim_truth rows, srm.split_summary and, when present, the A/A row of reports/validation.md.

CLI (make slide): python slide/build_slide.py [--db PATH] [--win-exp ID] [--harm-exp ID]
                  [--validation MD] [--out PATH]
"""

import argparse
import re
import sqlite3
import sys
from datetime import timedelta
from pathlib import Path

import plotly.graph_objects as go
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from abtest import boundaries, config, contracts, decision, srm, store
from abtest.contracts import ExperimentConfig, SimTruth, State

NAVY, GOLD, CRIMSON = "#003366", "#e8a517", "#C8102E"
INK, MUTED, RULE = "#1a202c", "#5a6472", "#d5dae1"
FONT = "Calibri, Carlito, 'Segoe UI', Helvetica, Arial, sans-serif"
PPT_FONT = "Calibri"
OUT = Path(__file__).resolve().parent / "one_slide.pptx"
W, H = 40 / 3, 7.5               # PowerPoint's 16:9 Widescreen: 12192000 x 6858000 EMU
CHART = (0.45, 1.95, 5.05, 3.6)  # x, y, w, h of the chart picture
WORKFLOW = ("Route", "Track", "Analyse", "Decide", "Roll out")
TITLE = "Slow & Steady: prompt changes ship on evidence, not judgement"
DATASETS = {"bttc": "historical Vani meeting calls", "buyer_logs": "historical buyer calls"}


# --- data -----------------------------------------------------------------------------------

def _pick(conn: sqlite3.Connection, state: State, scenario: str, exp_id: str | None) -> str | None:
    """exp_id if it is in `state`; else the newest run in `state`, the scenario's run first."""
    row = conn.execute(
        "SELECT experiment_id FROM experiments JOIN experiment_state USING (experiment_id) "
        "WHERE state = ? AND experiment_id = coalesce(?, experiment_id) "
        "ORDER BY scenario = ? DESC, created_ts_wall DESC", (state.value, exp_id, scenario)).fetchone()
    if row is None and exp_id:
        sys.exit(f"error: {exp_id} is not {state.value} in the store")
    return row and row[0]


def _run(conn: sqlite3.Connection, exp_id: str | None) -> dict | None:
    """Everything the slide says about one run, from its last look (the one that decided)."""
    if exp_id is None:
        return None
    cfg, st = store.get_config(conn, exp_id), store.get_state(conn, exp_id)
    looks = [d for d in store.read_decisions(conn, exp_id) if d.kind == "look"]
    last = looks[-1]
    days = (last.ts_sim - st.started_ts_sim) / timedelta(days=1)
    truth = conn.execute("SELECT payload FROM sim_truth WHERE experiment_id = ?", (exp_id,)).fetchone()
    scenario = conn.execute("SELECT scenario FROM experiments WHERE experiment_id = ?", (exp_id,)).fetchone()[0]
    return dict(id=exp_id, cfg=cfg, looks=looks, last=last, day=round(days), scenario=scenario,
                label=" / ".join(cfg.primary.dispositions),
                # stopped by the guardrail, not by the primary's harm boundary (rule 2's second half)
                guard_stop=st.state == State.STOPPED_HARM and not last.primary.z <= last.primary.harm_boundary,
                week=decision.weekly_impact(last.primary.lift_abs, last.n_A + last.n_B, days, cfg),
                truth=truth and SimTruth.model_validate_json(truth[0]))


def _design(cfg: ExperimentConfig) -> tuple:
    """What one chart draws for a run: its dataset and planned boundaries."""
    return cfg.dataset, boundaries.compute_boundaries(boundaries.planned_info_fractions(cfg), cfg)


def _aa_false_win(path: Path) -> tuple[str, str] | None:
    """(rate, runs) from the sequential A/A false-win row of reports/validation.md, if present."""
    try:
        m = re.search(r"\| false win, sequential[^|]*\| ([\d.]+%) \(\d+/(\d+)", path.read_text("utf-8"))
    except OSError:
        return None
    return m and (m[1], f"{int(m[2]):,}")


def _volume(scenario: str | None) -> str:
    try:
        scale = config.load_sim(config.load_scenario(scenario).sim).volume.scale
    except (OSError, ValueError, TypeError):
        return ""
    return "" if scale == 1 else f", volume ×{scale:g}"


def _guard(cfg: ExperimentConfig) -> str:
    s = cfg.secondary
    if s.type == "binary":
        return " / ".join(s.dispositions)
    return "call\u00a0duration" if s.column == "duration_s" else s.name


def _delta(x: float, cfg: ExperimentConfig) -> str:
    """A guardrail change in its units: relative % (continuous) or percentage points (binary)."""
    return f"{x * 100:+.2g}{'%' if cfg.secondary.type == 'continuous' else ' pp'}".replace("-", "−")


def _pct(x: float) -> str:
    return f"{x:+.0%}".replace("-", "−")


# --- chart ----------------------------------------------------------------------------------

def chart(win: dict | None, harm: dict | None, png: Path) -> Path:
    """Planned win/harm boundaries with both z-paths, direct-labelled. Layout is in 96-dpi px,
    the picture's size on the slide, so a 16 px font prints at 12 pt."""
    cfg = (win or harm)["cfg"]
    t = boundaries.planned_info_fractions(cfg)
    b = boundaries.compute_boundaries(t, cfg)
    harm_line = [-h for h in b.harm]
    zs = [d.primary.z for r in (win, harm) if r for d in r["looks"]]
    top, bottom = max([*b.win, *zs]) + 0.5, min(harm_line + zs) - 0.9
    fig = go.Figure()
    for edge, line, colour, fill, name, shift in (
            (top, b.win, GOLD, "rgba(232,165,23,0.13)", "Win boundary", -14),
            (bottom, harm_line, CRIMSON, "rgba(200,16,46,0.07)", "Harm boundary", 14)):
        fig.add_scatter(x=t, y=[edge] * len(t), mode="lines", line=dict(width=0), hoverinfo="skip")
        fig.add_scatter(x=t, y=line, mode="lines", fill="tonexty", fillcolor=fill,
                        line=dict(color=colour, width=2.5, dash="dash"))
        fig.add_annotation(x=t[-1], y=line[-1], text=name, showarrow=False, xanchor="right",
                           yshift=shift, font=dict(size=16, color=MUTED))
    for run, colour, shift in ((win, NAVY, 14), (harm, CRIMSON, -14)):
        if run is None:
            continue
        text = ("Winner promoted" if run is win else
                "Guardrail stop" if run["guard_stop"] else "Harmful stopped")
        x, z = [d.info_fraction for d in run["looks"]], [d.primary.z for d in run["looks"]]
        n = len(x) - 1  # looks before the deciding one: small markers, so only the decision stands out
        fig.add_scatter(x=x, y=z, mode="lines+markers", line=dict(color=colour, width=3.5),
                        marker=dict(size=[8] * n + [12], color=colour,
                                    line=dict(color="white", width=[1.5] * n + [2])))
        right = x[-1] > 0.6  # label to the left of the point, so it stays inside the plot
        fig.add_annotation(x=x[-1], y=z[-1], text=f"<b>{text}, day {run['day']}</b>", showarrow=False,
                           xanchor="right" if right else "left", xshift=-8 if right else 8, yshift=shift,
                           font=dict(size=16, color=INK))
        held = next((d for d in run["looks"][:-1] if d.primary.z >= d.primary.win_boundary), None)
        if held:  # rule 4: the primary crossed but the guardrail had not cleared yet
            fig.add_annotation(x=held.info_fraction, y=held.primary.z, ax=0, ay=52, arrowhead=0,
                               arrowcolor=MUTED, text=f"look {held.look}: crossed, guardrail<br>not yet clear",
                               font=dict(size=16, color=MUTED))
    fig.add_hline(y=0, line=dict(color=RULE, width=1.5))
    fig.update_layout(width=round(CHART[2] * 96), height=round(CHART[3] * 96), showlegend=False,
                      paper_bgcolor="white", plot_bgcolor="white", margin=dict(l=56, r=12, t=8, b=52),
                      font=dict(family=FONT, size=16, color=INK))
    fig.update_xaxes(range=[0.17, 1.02], tickvals=[0.2, 0.4, 0.6, 0.8, 1.0], tickformat=".0%",
                     showgrid=False, linecolor=RULE, title=dict(
                         text="Share of the planned sample (information fraction)",
                         font=dict(size=16, color=MUTED), standoff=6))
    fig.update_yaxes(range=[bottom, top], gridcolor="#eef0f3", zeroline=False, linecolor=RULE,
                     title=dict(text="z-statistic, B vs A", font=dict(size=16, color=MUTED), standoff=4))
    png.parent.mkdir(parents=True, exist_ok=True)
    fig.write_image(png, scale=4)
    return png


# --- pptx -----------------------------------------------------------------------------------

def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


def _text(shape_or_box, paras, size: float, color: str = INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, space: float = 0, margin: float = 0):
    """Fill a text frame. paras: list of paragraphs, each a str or a list of runs
    (text, bold, colour[, size])."""
    tf = shape_or_box.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment, p.space_before = align, Pt(space if i else 0)
        for text, bold, colour, *pt in ([(para, False, color)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = text
            r.font.size, r.font.bold, r.font.name = Pt(pt[0] if pt else size), bold, PPT_FONT
            r.font.color.rgb = _rgb(colour)
    return shape_or_box


def _box(slide, x, y, w, h):
    return slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))


def _shape(slide, kind, x, y, w, h, fill: str):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _header(slide, x, w, text):
    _text(_box(slide, x, 1.52, w, 0.3), [[(text.upper(), True, NAVY)]], 13, anchor=MSO_ANCHOR.BOTTOM)
    _shape(slide, MSO_SHAPE.RECTANGLE, x, 1.85, w, 0.03, GOLD)


def build(db: Path, out: Path = OUT, png: Path | None = None, win_exp: str | None = None,
          harm_exp: str | None = None, validation: Path | None = None) -> Path:
    if not Path(db).is_file():
        sys.exit(f"error: no store at {db}; run make demo first, or pass --db")
    conn = sqlite3.connect(f"file:{db}?mode=ro", uri=True)  # the slide never writes the store
    conn.row_factory = sqlite3.Row
    try:
        win = _run(conn, _pick(conn, State.PROMOTED, "win", win_exp))
        harm = _run(conn, _pick(conn, State.STOPPED_HARM, "harm", harm_exp))
        if not (win or harm):
            sys.exit(f"error: {db} has no PROMOTED or STOPPED_HARM experiment; run make demo first")
        if win and harm and _design(win["cfg"]) != _design(harm["cfg"]):
            sys.exit(f"error: {win['id']} and {harm['id']} differ in dataset or look plan, so one chart "
                     "cannot show both; pass a --win-exp and --harm-exp of the same experiment")
        ref = win or harm
        split = srm.split_summary(conn, ref["id"])
    finally:
        conn.close()
    cfg = ref["cfg"]
    guard = _guard(cfg)
    aa = _aa_false_win(validation or contracts.REPORTS_DIR / "validation.md")
    png = chart(win, harm, png or contracts.REPORTS_DIR / "slide_zpath.png")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Headline band: the claim, then the result in business terms
    said = []
    if win:
        said.append(f"Winner promoted on day {win['day']}: ≈ {win['week']:+,.0f} {win['label']} a week.")
    if harm:
        said.append(f"Harmful variant stopped on day {harm['day']}" + (
            f" by its {_guard(harm['cfg'])} guardrail." if harm["guard_stop"] else
            f": ≈ {abs(harm['week']):,.0f} {'' if win else harm['label'] + ' '}a week protected."))
    _shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, 1.28, NAVY)
    _shape(slide, MSO_SHAPE.RECTANGLE, 0, 1.28, W, 0.06, GOLD)
    _text(_box(slide, 0.45, 0.18, W - 0.9, 0.62), [[(TITLE, True, "#ffffff")]], 30, anchor=MSO_ANCHOR.MIDDLE)
    _text(_box(slide, 0.45, 0.8, W - 0.9, 0.4), [[(" ".join(said), False, "#f3d58f")]], 16,
          anchor=MSO_ANCHOR.MIDDLE)

    # Left: the result, business first
    x, w = CHART[0], CHART[2]
    _header(slide, x, w, "Result: two simulated A/B runs")
    slide.shapes.add_picture(str(png), Inches(x), Inches(CHART[1]), Inches(w), Inches(CHART[3]))
    tiles = []
    if win:
        tiles.append((NAVY, f"≈ {win['week']:+,.0f} / week", f"{win['label']} gained",
                      f"Promoted on day {win['day']} (look {win['last'].look})"))
    if harm:
        hc = harm["cfg"]
        big, what = (f"≈ {abs(harm['week']):,.0f} / week", f"{harm['label']} protected")
        if harm["guard_stop"]:
            big, what = _delta(harm["last"].guardrail.delta, hc), f"Rise in {_guard(hc)} avoided"
        tiles.append((CRIMSON, big, what, f"Stopped on day {harm['day']} (look {harm['last'].look})"))
    tw = (w - 0.25 * (len(tiles) - 1)) / len(tiles)
    for i, (colour, big, what, when) in enumerate(tiles):
        tx = x + i * (tw + 0.25)
        _shape(slide, MSO_SHAPE.RECTANGLE, tx, 5.75, 0.07, 0.9, colour)
        _text(_box(slide, tx + 0.18, 5.7, tw - 0.18, 1.0),
              [[(big, True, colour, 24)], [(what, True, INK)], [(when, False, MUTED, 12)]], 13)

    # Middle: method (bullets), then the decision rules, numbered in the engine's order (§2.8)
    x, w = 5.8, 4.15
    _header(slide, x, w, "Method and decision rules")
    s = cfg.secondary
    margin = _delta(s.margin if s.worse_if == "increase" else -s.margin, cfg)
    lines = [
        (None, "Unit:", " the lead, from its first answered call; stats per lead, never per call."),
        (None, "Alpha spending:", f" O'Brien-Fleming win (α {cfg.alpha_win:.1%}), no lucky early wins; "
                                  f"Pocock harm (α {cfg.alpha_harm:.0%}), fast stops; any number of looks."),
        ("1", "SRM halts:", f" split off its {cfg.allocation.B:.0%} (χ² p < {cfg.srm_p_threshold:g}): "
                            "results untrusted, all traffic to A."),
        ("2", "Stop early:", f" z ≤ harm boundary, or the {guard} guardrail's repeated CI fully beyond "
                             f"{margin}."),
        ("3", "Promote:", f" z ≥ win boundary and the guardrail's repeated CI inside {margin}."),
    ]
    slot, gap = 0.86, 0.36  # equal slots, text centred in each, so line counts never misalign
    _text(_box(slide, x, 2.0 + 2 * slot + 0.04, w, gap - 0.08),
          [[("At every look, in this order; the first match decides:", True, MUTED)]], 12,
          anchor=MSO_ANCHOR.MIDDLE)
    for i, (num, lead, rest) in enumerate(lines):
        y = 2.0 + i * slot + (gap if num else 0)
        if num:
            badge = _shape(slide, MSO_SHAPE.OVAL, x, y + slot / 2 - 0.18, 0.36, 0.36, GOLD)
            _text(badge, [[(num, True, NAVY)]], 14, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            _shape(slide, MSO_SHAPE.OVAL, x + 0.11, y + slot / 2 - 0.07, 0.14, 0.14, NAVY)
        _text(_box(slide, x + 0.5, y + 0.03, w - 0.5, slot - 0.06),
              [[(lead, True, NAVY), (rest, False, INK)]], 13.5, anchor=MSO_ANCHOR.MIDDLE)

    # Right: the workflow as shapes
    x, w = 10.25, 2.63
    _header(slide, x, w, "Workflow")
    a = split.assigned
    captions = (
        [f"Sticky hash buckets, {cfg.allocation.A:.0%}/{cfg.allocation.B:.0%}",
         f"Achieved B {a.observed_share_B:.1%}, {split.leads_in_both_arms} in both arms"
         if a.n_A + a.n_B else "No exposures logged"],
        ["Exposures and calls per lead"],
        ["z, CIs, repeated CIs per look"],
        ["First match: SRM, harm, win", "Else, at the last look: A stays"],
        ["Live allocation + audit log"],
    )
    bh, gap = 0.79, 0.2
    for i, (name, caps) in enumerate(zip(WORKFLOW, captions)):
        y = 2.05 + i * (bh + gap)
        last = i == len(WORKFLOW) - 1
        box = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, bh, GOLD if last else NAVY)
        box.adjustments[0] = 0.12
        ink = INK if last else "#dce5ef"
        _text(box, [[(name, True, NAVY if last else "#ffffff", 15)], *[[(c, False, ink)] for c in caps]],
              12, anchor=MSO_ANCHOR.MIDDLE, margin=0.1)
        if not last:
            _shape(slide, MSO_SHAPE.DOWN_ARROW, x + w / 2 - 0.12, y + bh + 0.025, 0.24, gap - 0.05, GOLD)

    # Footer: disclosure and sources
    _shape(slide, MSO_SHAPE.RECTANGLE, 0.45, 6.9, W - 0.9, 0.015, RULE)
    # the realised true lift (what B really did in this run), not the injected one
    truth = [(_pct(r["truth"].realised_primary_rel_lift), _pct(r["last"].primary.lift_rel))
             for r in (win, harm) if r and r["truth"]]
    true, est = (" / ".join(x) for x in zip(*truth)) if truth else ("", "")
    _text(_box(slide, 0.45, 6.94, W - 0.9, 0.5), [
        f"Simulated traffic from {DATASETS.get(cfg.dataset, cfg.dataset)}{_volume(ref['scenario'])}, "
        "replayed with injected effects the engine never sees" + (
            f" (true lift {true}; engine estimated {est})." if truth else "."),
        (f"A/A check: {aa[0]} false wins in {aa[1]} no-effect runs. " if aa else "")
        + "Weekly figures: engine estimate at 100% rollout. Sources: the store's decision log"
        + (", reports/validation.md." if aa else ".")],
          12, MUTED)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(prog="python slide/build_slide.py", description=__doc__.split("\n")[0])
    ap.add_argument("--db", type=Path, default=None, help="store to read (default: contracts.DB_PATH)")
    ap.add_argument("--win-exp", help="a PROMOTED experiment id (default: the win scenario's run)")
    ap.add_argument("--harm-exp", help="a STOPPED_HARM experiment id of the same design "
                                       "(default: the harm scenario's run)")
    ap.add_argument("--validation", type=Path, default=None,
                    help="validation.md to quote the A/A rate from (default: reports/validation.md)")
    ap.add_argument("--out", type=Path, default=OUT)
    args = ap.parse_args(argv)
    out = build(args.db or contracts.DB_PATH, args.out, win_exp=args.win_exp, harm_exp=args.harm_exp,
                validation=args.validation)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()

