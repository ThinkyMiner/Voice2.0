"""slide/build_slide.py on a small synthetic store (no data/ needed)."""

import importlib.util
from datetime import datetime, timedelta

import pandas as pd
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from abtest import contracts, decision, router, store
from abtest.contracts import (Boundaries, ExperimentConfig, FrozenConfig, LookResult, PrimaryStats,
                               SecondaryStats, SrmResult, State, Verdict)

_spec = importlib.util.spec_from_file_location("build_slide", contracts.ROOT / "slide" / "build_slide.py")
slide_mod = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(slide_mod)

T0 = datetime(2026, 10, 1)


def _cfg(exp_id: str, looks: int = 7) -> ExperimentConfig:
    return ExperimentConfig.model_validate({
        "id": exp_id, "dataset": "bttc", "unit": "lead_id", "trigger": "first_answered_call",
        "salt": "s", "allocation": {"A": 0.9, "B": 0.1},
        "variants": {"A": {"base": "prompts/base_buyer_vani.txt"}, "B": {"base": "prompts/base_buyer_vani.txt"}},
        "window": {"days": 28, "looks": looks}, "maturation_days": 2, "baseline_p": 0.18,
        "primary": {"name": "m", "type": "binary", "dispositions": ["Meeting Fixed"], "better": "increase",
                    "mde_rel": 0.15},
        "secondary": {"name": "dur", "type": "continuous", "role": "guardrail", "worse_if": "increase",
                      "margin_rel": 0.1},
        "alpha_win": 0.025, "alpha_harm": 0.05, "power": 0.8, "min_info_first_look": 0.2,
        "srm_p_threshold": 0.001, "promotion": "auto",
    })


def _look(exp_id: str, k: int, lift: float, z: float) -> LookResult:
    """6,000 matured leads, look k on day 7 + 3k."""
    p = PrimaryStats(rate_A=0.18, rate_B=0.18 + lift, lift_abs=lift, lift_rel=lift / 0.18,
                     ci95=(lift - 0.03, lift + 0.03), se_pooled=0.015, se_unpooled=0.016, z=z,
                     detectable_mde_rel=0.2)
    s = SecondaryStats(name="dur", type="continuous", role="guardrail", mean_A=40, mean_B=40.4,
                       rel_change=0.01, delta=0.01, delta_type="relative", se=0.02, ci95=(-0.03, 0.05),
                       rci=(-0.05, 0.07), margin=0.1)
    srm = SrmResult(n_A=9000, n_B=1000, configured_share_B=0.1, observed_share_B=0.1,
                    ci95=(0.094, 0.106), chi2=0.0, p=1.0)
    return LookResult(experiment_id=exp_id, look=k, ts_sim=T0 + timedelta(days=7 + 3 * k),
                      is_final=False, info_fraction=0.2 * k, n_A=5400, n_B=600, primary=p, secondary=s,
                      srm_assigned=srm, srm_triggered=srm)


def _run(conn, exp_id: str, steps: list[tuple[float, float, str]], looks: int = 7) -> None:
    cfg = _cfg(exp_id, looks)
    frozen = FrozenConfig(config=cfg, canonical_json=cfg.model_dump_json(), config_hash="h",
                          prompt_hash={"A": "a", "B": "b"})
    store.create_experiment(conn, frozen)
    state = decision.start(conn, frozen, T0)
    for k, (lift, z, verdict) in enumerate(steps, 1):
        bounds = Boundaries(info_fractions=tuple(0.2 * i for i in range(1, k + 1)), win=(3.0,) * k,
                            harm=(2.2,) * k)
        state = decision.apply_verdict(conn, state, _look(exp_id, k, lift, z), bounds,
                                       Verdict(decision=verdict, rule=0 if verdict == "CONTINUE" else 3,
                                               reason="Rule fired."))


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    """Built once: a promoted run (2 looks), a harm-stopped run (1 look), 400 exposures."""
    tmp_path = tmp_path_factory.mktemp("slide")
    db = tmp_path / "s.db"
    conn = store.connect(db)
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(decision, "JSONL", None)  # never touch the repo's logs/decisions.jsonl
        _run(conn, "win-exp", [(0.03, 2.0, "CONTINUE"), (0.044, 3.4, State.PROMOTED)])
        _run(conn, "harm-exp", [(-0.05, -2.6, State.STOPPED_HARM)])
    leads = [f"l{i}" for i in range(400)]
    buckets = [router.bucket("s", lead) for lead in leads]
    store.insert_df(conn, "exposures", pd.DataFrame({
        "experiment_id": "win-exp", "lead_id": leads, "bucket": buckets, "ts": T0 + timedelta(hours=1),
        "variant": [router.variant_for(b, _cfg("x").allocation) for b in buckets]}))
    conn.close()
    validation = tmp_path / "validation.md"
    validation.write_text("| false win, sequential (PROMOTED or PENDING_APPROVAL) | 2.30% (46/2000; "
                          "95% CI 1.69% to 3.06%) | 2.5% (nominal) |\n", encoding="utf-8")
    try:
        out = slide_mod.build(db, tmp_path / "deck.pptx", png=tmp_path / "z.png", validation=validation)
    except RuntimeError as e:  # kaleido renders the chart PNG with a headless Chrome
        if "Chrome" not in str(e):
            raise
        pytest.skip(f"kaleido cannot start Chrome: {e}")
    return Presentation(out), tmp_path


def _texts(slide) -> list[str]:
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def test_one_16_9_slide_with_title_chart_and_workflow(deck):
    prs, tmp = deck
    assert len(prs.slides) == 1
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)  # PowerPoint's Widescreen
    slide = prs.slides[0]
    texts = _texts(slide)
    assert slide_mod.TITLE in texts
    assert any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
    assert (tmp / "z.png").read_bytes()[:4] == b"\x89PNG"
    steps = sorted((s.top, s.text_frame.paragraphs[0].text) for s in slide.shapes
                   if s.has_text_frame and s.text_frame.paragraphs[0].text in slide_mod.WORKFLOW)
    assert [name for _, name in steps] == list(slide_mod.WORKFLOW)


def test_every_shape_inside_the_slide(deck):
    prs, _ = deck
    for s in prs.slides[0].shapes:
        assert s.left >= 0 and s.top >= 0, s.name
        assert s.left + s.width <= prs.slide_width and s.top + s.height <= prs.slide_height, s.name


def test_numbers_come_from_the_store(deck):
    prs, _ = deck
    text = "\n".join(_texts(prs.slides[0]))
    cfg = _cfg("win-exp")
    win = decision.weekly_impact(0.044, 6000, 13, cfg)   # promoted at look 2, day 13
    harm = decision.weekly_impact(-0.05, 6000, 10, cfg)  # stopped at look 1, day 10
    assert f"≈ {win:+,.0f} / week" in text and "Promoted on day 13 (look 2)" in text
    assert f"≈ {abs(harm):,.0f} / week" in text and "Stopped on day 10 (look 1)" in text
    assert f"Winner promoted on day 13: ≈ {win:+,.0f} Meeting Fixed a week." in text  # the headline
    assert "A/A check: 2.30% false wins in 2,000 no-effect runs" in text
    assert "Achieved B" in text and "0 in both arms" in text


def test_rules_in_the_engines_order(deck):
    """PLAN §2.8: SRM, then harm, then promote; the badges number them 1-3 in that order."""
    prs, _ = deck
    rules = sorted((s.top, s.text_frame.text.split(":")[0]) for s in prs.slides[0].shapes if s.has_text_frame
                   and s.text_frame.text.startswith(("SRM halts:", "Stop early:", "Promote:")))
    assert [name for _, name in rules] == ["SRM halts", "Stop early", "Promote"]


def test_guardrail_stop_and_mismatched_designs(tmp_path, monkeypatch):
    """A harm stop with z above the harm boundary is a guardrail stop; runs of different look
    plans cannot share one chart, so the build refuses them (before any chart is drawn)."""
    monkeypatch.setattr(decision, "JSONL", None)
    db = tmp_path / "g.db"
    conn = store.connect(db)
    _run(conn, "harm-exp", [(-0.05, -2.6, State.STOPPED_HARM)])
    _run(conn, "guard-exp", [(0.01, 0.8, State.STOPPED_HARM)])
    _run(conn, "win5-exp", [(0.044, 3.4, State.PROMOTED)], looks=5)
    assert slide_mod._run(conn, "guard-exp")["guard_stop"]
    assert not slide_mod._run(conn, "harm-exp")["guard_stop"]
    conn.close()
    with pytest.raises(SystemExit, match="differ in dataset or look plan"):
        slide_mod.build(db, tmp_path / "x.pptx", png=tmp_path / "x.png")

